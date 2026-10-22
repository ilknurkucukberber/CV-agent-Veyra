from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

PDF_TYPES = {"application/pdf"}
DOCX_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
IMAGE_TYPES = {"image/jpeg", "image/png"}

ALLOWED_CONTENT_TYPES = PDF_TYPES | DOCX_TYPES | IMAGE_TYPES


@dataclass
class ExtractionResult:
    text: str
    page_count: int | None = None
    table_count: int = 0
    multi_column_pages: list[int] = field(default_factory=list)


def extract_text(file_path: Path, content_type: str) -> ExtractionResult:
    if content_type in PDF_TYPES:
        return _extract_pdf(file_path)
    if content_type in DOCX_TYPES:
        return _extract_docx(file_path)
    if content_type in IMAGE_TYPES:
        return _extract_image_ocr(file_path)
    raise ValueError(f"Desteklenmeyen dosya türü: {content_type}")


def _extract_pdf(file_path: Path) -> ExtractionResult:
    texts: list[str] = []
    table_count = 0
    multi_column_pages: list[int] = []

    with fitz.open(file_path) as doc:
        for page_index, page in enumerate(doc):
            texts.append(page.get_text())
            try:
                table_count += len(page.find_tables().tables)
            except Exception:
                pass  # tablo tespiti best-effort; başarısızlığı analizi engellemesin
            if _looks_multi_column(page):
                multi_column_pages.append(page_index + 1)

        return ExtractionResult(
            text="\n".join(texts).strip(),
            page_count=doc.page_count,
            table_count=table_count,
            multi_column_pages=multi_column_pages,
        )


def _looks_multi_column(page: "fitz.Page") -> bool:
    """Heuristik kolon tespiti: metin bloklarının önemli bir kısmı sayfanın sağ
    yarısından başlıyorsa çok kolonlu düzen olasıdır (ATS'leri zorlayabilir)."""
    blocks = [b for b in page.get_text("blocks") if b[6] == 0 and (b[4] or "").strip()]
    if len(blocks) < 4:
        return False
    midpoint = page.rect.width / 2
    right_start = sum(1 for b in blocks if b[0] > midpoint)
    return right_start / len(blocks) > 0.3


def _extract_docx(file_path: Path) -> ExtractionResult:
    doc = Document(str(file_path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ExtractionResult(
        text="\n".join(parts).strip(),
        page_count=None,  # docx'te güvenilir sayfa sayısı render olmadan alınamaz
        table_count=len(doc.tables),
    )


def _extract_image_ocr(file_path: Path) -> ExtractionResult:
    """Görsel CV'ler için OCR kullanarak metin çıkarır."""
    import pytesseract
    from PIL import Image

    try:
        image = Image.open(str(file_path))
        # Türkçe ve İngilizce dillerini destekleyecek şekilde ayarlanır
        text = pytesseract.image_to_string(image, lang="tur+eng")
        return ExtractionResult(
            text=text.strip(),
            page_count=1,
            table_count=0,
        )
    except Exception as e:
        raise ValueError(f"Görsel üzerinden OCR işlemi başarısız oldu: {e}")
